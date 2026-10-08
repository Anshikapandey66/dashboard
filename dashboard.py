import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# --------------------------------------------------
# 1. EXPENSE DATA
# --------------------------------------------------

data = {
    "Date": [
        "2026-10-01", "2026-10-02", "2026-10-03",
        "2026-10-04", "2026-10-05", "2026-10-06",
        "2026-10-07", "2026-10-08", "2026-10-09",
        "2026-10-10", "2026-10-11", "2026-10-12",
        "2026-10-13", "2026-10-14", "2026-10-15",
        "2026-10-16", "2026-10-17", "2026-10-18",
        "2026-10-19", "2026-10-20"
    ],

    "Category": [
        "Food", "Travel", "Shopping", "Food",
        "Education", "Bills", "Travel", "Food",
        "Shopping", "Education", "Food", "Travel",
        "Bills", "Food", "Shopping", "Travel",
        "Education", "Food", "Bills", "Shopping"
    ],

    "Description": [
        "Lunch", "Auto", "Clothes", "Dinner",
        "Books", "Electricity", "Metro", "Breakfast",
        "Shoes", "Course", "Lunch", "Cab",
        "Internet", "Dinner", "Bag", "Auto",
        "Stationery", "Lunch", "Mobile Recharge", "T-Shirt"
    ],

    "Amount": [
        180, 120, 850, 250, 600,
        1200, 80, 120, 1500, 900,
        200, 350, 700, 300, 950,
        100, 250, 220, 399, 650
    ],

    "Payment_Method": [
        "UPI", "Cash", "Card", "UPI", "UPI",
        "UPI", "Card", "Cash", "Card", "UPI",
        "UPI", "UPI", "UPI", "Cash", "Card",
        "Cash", "UPI", "UPI", "UPI", "Card"
    ]
}

df = pd.DataFrame(data)

df["Date"] = pd.to_datetime(df["Date"])

# --------------------------------------------------
# 2. CREATE EXCEL WORKBOOK
# --------------------------------------------------

wb = Workbook()

ws_data = wb.active
ws_data.title = "Expense Data"

ws_dashboard = wb.create_sheet("Dashboard")
ws_analysis = wb.create_sheet("Analysis")

# --------------------------------------------------
# 3. EXPENSE DATA SHEET
# --------------------------------------------------

headers = list(df.columns)

for col_num, header in enumerate(headers, 1):
    cell = ws_data.cell(row=1, column=col_num)
    cell.value = header
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

for row_num, row in enumerate(df.itertuples(index=False), 2):

    for col_num, value in enumerate(row, 1):

        cell = ws_data.cell(
            row=row_num,
            column=col_num,
            value=value
        )

        if col_num == 1:
            cell.number_format = "DD-MM-YYYY"

        if col_num == 4:
            cell.number_format = '₹#,##0'

# Create Excel Table

table_ref = f"A1:E{len(df) + 1}"

tab = Table(
    displayName="ExpenseTable",
    ref=table_ref
)

style = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)

tab.tableStyleInfo = style

ws_data.add_table(tab)

# Column widths

widths = {
    "A": 15,
    "B": 18,
    "C": 25,
    "D": 15,
    "E": 20
}

for column, width in widths.items():
    ws_data.column_dimensions[column].width = width

ws_data.freeze_panes = "A2"

# --------------------------------------------------
# 4. ANALYSIS SHEET
# --------------------------------------------------

ws_analysis["A1"] = "Expense Analysis"

ws_analysis["A1"].font = Font(
    bold=True,
    size=18
)

# KPI calculations

ws_analysis["A3"] = "Total Expense"
ws_analysis["B3"] = "=SUM('Expense Data'!D2:D21)"

ws_analysis["A4"] = "Average Expense"
ws_analysis["B4"] = "=AVERAGE('Expense Data'!D2:D21)"

ws_analysis["A5"] = "Highest Expense"
ws_analysis["B5"] = "=MAX('Expense Data'!D2:D21)"

ws_analysis["A6"] = "Number of Expenses"
ws_analysis["B6"] = "=COUNT('Expense Data'!D2:D21)"

for row in range(3, 7):
    ws_analysis[f"A{row}"].font = Font(bold=True)
    ws_analysis[f"B{row}"].number_format = '₹#,##0'

# Category analysis

ws_analysis["A9"] = "Category"
ws_analysis["B9"] = "Total Expense"

categories = [
    "Food",
    "Travel",
    "Shopping",
    "Education",
    "Bills"
]

for i, category in enumerate(categories, 10):

    ws_analysis[f"A{i}"] = category

    ws_analysis[f"B{i}"] = (
        f'=SUMIF(\'Expense Data\'!B:B,A{i},'
        f'\'Expense Data\'!D:D)'
    )

    ws_analysis[f"B{i}"].number_format = '₹#,##0'

# Payment analysis

ws_analysis["D9"] = "Payment Method"
ws_analysis["E9"] = "Total Expense"

payments = [
    "UPI",
    "Cash",
    "Card"
]

for i, payment in enumerate(payments, 10):

    ws_analysis[f"D{i}"] = payment

    ws_analysis[f"E{i}"] = (
        f'=SUMIF(\'Expense Data\'!E:E,D{i},'
        f'\'Expense Data\'!D:D)'
    )

    ws_analysis[f"E{i}"].number_format = '₹#,##0'

# --------------------------------------------------
# 5. DASHBOARD
# --------------------------------------------------

ws_dashboard.sheet_view.showGridLines = False

ws_dashboard.merge_cells("B2:I3")

ws_dashboard["B2"] = "EXPENSE TRACKER DASHBOARD"

ws_dashboard["B2"].font = Font(
    bold=True,
    size=22
)

ws_dashboard["B2"].alignment = Alignment(
    horizontal="center",
    vertical="center"
)

# KPI Cards

cards = [
    ("B5:C7", "TOTAL EXPENSE", "='Analysis'!B3"),
    ("D5:E7", "AVERAGE EXPENSE", "='Analysis'!B4"),
    ("F5:G7", "HIGHEST EXPENSE", "='Analysis'!B5"),
    ("H5:I7", "TOTAL RECORDS", "='Analysis'!B6")
]

for cell_range, title, formula in cards:

    ws_dashboard.merge_cells(cell_range)

    start_cell = cell_range.split(":")[0]

    cell = ws_dashboard[start_cell]

    cell.value = f"{title}\n{formula}"

    cell.font = Font(
        bold=True,
        size=13
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    cell.border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

# --------------------------------------------------
# 6. CATEGORY PIE CHART
# --------------------------------------------------

pie = PieChart()

labels = Reference(
    ws_analysis,
    min_col=1,
    min_row=10,
    max_row=14
)

values = Reference(
    ws_analysis,
    min_col=2,
    min_row=9,
    max_row=14
)

pie.add_data(
    values,
    titles_from_data=True
)

pie.set_categories(labels)

pie.title = "Expense by Category"

pie.height = 8
pie.width = 11

ws_dashboard.add_chart(
    pie,
    "B9"
)

# --------------------------------------------------
# 7. PAYMENT METHOD BAR CHART
# --------------------------------------------------

bar = BarChart()

values = Reference(
    ws_analysis,
    min_col=5,
    min_row=9,
    max_row=12
)

labels = Reference(
    ws_analysis,
    min_col=4,
    min_row=10,
    max_row=12
)

bar.add_data(
    values,
    titles_from_data=True
)

bar.set_categories(labels)

bar.title = "Expense by Payment Method"

bar.y_axis.title = "Amount"

bar.x_axis.title = "Payment Method"

bar.height = 8
bar.width = 11

ws_dashboard.add_chart(
    bar,
    "H9"
)

# --------------------------------------------------
# 8. DASHBOARD COLUMN WIDTHS
# --------------------------------------------------

for column in range(1, 11):
    ws_dashboard.column_dimensions[
        get_column_letter(column)
    ].width = 15

# --------------------------------------------------
# 9. SAVE FILE
# --------------------------------------------------

file_name = "Expense_Tracker.xlsx"

wb.save(file_name)

print("--------------------------------")
print("Expense Tracker Created!")
print("--------------------------------")
print(f"File: {file_name}")
print("Dashboard created successfully.")
